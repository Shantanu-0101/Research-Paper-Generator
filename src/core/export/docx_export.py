import io
import re
import os
from docx import Document
from docx.shared import Inches

def generate_docx(title: str, sections: list[tuple[str, str]]) -> bytes:
    doc = Document()
    
    # Add Title
    title_para = doc.add_heading(title, 0)
    title_para.alignment = 1  # Center alignment
    
    # Regex to find our image placeholders: ![caption](image_path)
    img_pattern = re.compile(r"!\[(.*?)\]\((.*?)\)")

    for section_title, content in sections:
        doc.add_heading(section_title, level=1)
        
        for line in content.split("\n"):
            line = line.strip()
            if not line:
                continue
                
            # Parse Markdown Headings (which we forced the LLM to use!)
            if line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            else:
                # Check if the line is an image
                match = img_pattern.search(line)
                if match:
                    caption, img_path = match.groups()
                    if os.path.exists(img_path):
                        try:
                            doc.add_picture(img_path, width=Inches(5.5))
                            cap_para = doc.add_paragraph(caption)
                            cap_para.alignment = 1 # Center align the caption
                        except Exception:
                            doc.add_paragraph(f"[Image could not be loaded: {img_path}]")
                    else:
                        doc.add_paragraph(f"[Image not found locally: {img_path}]")
                else:
                    # Just normal paragraph text
                    doc.add_paragraph(line)
                    
        doc.add_page_break()

    # Save to memory buffer instead of disk
    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()